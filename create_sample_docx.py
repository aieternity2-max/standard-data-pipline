from docx import Document


document = Document()

document.add_heading(
    "Customer Data Report",
    level=1
)

document.add_paragraph(
    "Customer 101: Alice"
)

document.add_paragraph(
    "Customer 102: Bob"
)

document.add_paragraph(
    "Customer 103: Charlie"
)

document.add_paragraph(
    "All customers have completed their registration."
)

document.save(
    "tests/fixtures/sample.docx"
)

print("sample.docx created successfully")
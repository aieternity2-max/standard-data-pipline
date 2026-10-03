from pathlib import Path

from docx import Document as DocxDocument

from app.loaders.base import BaseLoader
from app.models.document import Document


class DOCXLoader(BaseLoader):
    """
    Loader for DOCX files.

    Supports:
    1. Normal paragraph loading using load()
    2. Batch loading using load_batches()
    """

    def load(self, source: str | Path) -> list[Document]:
        """
        Load paragraphs from a DOCX file.

        Each non-empty paragraph becomes a Document.
        """

        file_path = Path(source)

        self._validate_file(file_path)

        try:
            docx = DocxDocument(str(file_path))

        except Exception as exc:
            raise ValueError(
                f"Unable to read DOCX file: {file_path}"
            ) from exc

        documents = []

        paragraph_number = 0

        for paragraph in docx.paragraphs:

            content = paragraph.text.strip()

            if not content:
                continue

            paragraph_number += 1

            document = Document(
                id=f"{file_path.stem}-{paragraph_number}",
                source=str(file_path),
                source_type="docx",
                content=content,
                metadata={
                    "file_name": file_path.name,
                    "paragraph_number": paragraph_number,
                },
            )

            documents.append(document)

        return documents

    def load_batches(
        self,
        source: str | Path,
        batch_size: int = 50,
    ):
        """
        Load DOCX paragraphs in batches.

        Each batch contains multiple paragraph Documents.
        """

        file_path = Path(source)

        self._validate_file(file_path)

        if batch_size <= 0:
            raise ValueError(
                "batch_size must be greater than 0"
            )

        try:
            docx = DocxDocument(str(file_path))

        except Exception as exc:
            raise ValueError(
                f"Unable to read DOCX file: {file_path}"
            ) from exc

        batch = []

        paragraph_number = 0

        for paragraph in docx.paragraphs:

            content = paragraph.text.strip()

            if not content:
                continue

            paragraph_number += 1

            document = Document(
                id=f"{file_path.stem}-{paragraph_number}",
                source=str(file_path),
                source_type="docx",
                content=content,
                metadata={
                    "file_name": file_path.name,
                    "paragraph_number": paragraph_number,
                },
            )

            batch.append(document)

            if len(batch) == batch_size:
                yield batch
                batch = []

        if batch:
            yield batch

    @staticmethod
    def _validate_file(file_path: Path) -> None:
        """
        Validate DOCX input.
        """

        if not file_path.exists():
            raise FileNotFoundError(
                f"DOCX file not found: {file_path}"
            )

        if not file_path.is_file():
            raise ValueError(
                f"Source is not a file: {file_path}"
            )

        if file_path.suffix.lower() != ".docx":
            raise ValueError(
                f"Unsupported DOCX file type: {file_path.suffix}"
            )